"""
Simple PDF converter using fpdf2 to convert Word content to PDF
Alternative approach since ReportLab has compatibility issues
"""
from fpdf import FPDF
from docx import Document
import os

class PDF(FPDF):
    def header(self):
        self.set_font('Arial', 'B', 12)
        self.cell(0, 10, '', 0, 1, 'C')

    def footer(self):
        self.set_y(-15)
        self.set_font('Arial', 'I', 8)
        self.cell(0, 10, f'Page {self.page_no()}', 0, 0, 'C')

    def chapter_title(self, title):
        self.set_font('Arial', 'B', 14)
        self.cell(0, 10, title, 0, 1, 'L')
        self.ln(4)

    def chapter_body(self, body):
        self.set_font('Arial', '', 11)
        self.multi_cell(0, 6, body)
        self.ln()

def convert_docx_to_pdf(docx_path, pdf_path):
    """
    Convert Word document to PDF by extracting text
    """
    # Read Word document
    doc = Document(docx_path)
    
    # Create PDF
    pdf = PDF()
    pdf.add_page()
    
    # Title page
    pdf.set_font('Arial', 'B', 20)
    pdf.cell(0, 20, '', 0, 1)
    pdf.cell(0, 10, 'INCIDENT INVESTIGATION REPORT', 0, 1, 'C')
    pdf.set_font('Arial', 'B', 16)
    pdf.cell(0, 10, 'BP Texas City Refinery Explosion', 0, 1, 'C')
    pdf.set_font('Arial', '', 12)
    pdf.cell(0, 10, 'March 23, 2005', 0, 1, 'C')
    pdf.ln(20)
    
    # Extract and add content from Word document
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
            
        # Check if it's a heading (based on style or formatting)
        if paragraph.style.name.startswith('Heading'):
            pdf.add_page()
            pdf.chapter_title(text)
        else:
            # Encode text to latin-1, replacing unsupported characters
            try:
                safe_text = text.encode('latin-1', 'replace').decode('latin-1')
                pdf.chapter_body(safe_text)
            except Exception as e:
                print(f"Warning: Could not add paragraph: {str(e)[:50]}")
                continue
    
    # Save PDF
    pdf.output(pdf_path)
    print(f"✅ PDF created successfully: {pdf_path}")
    
    # Get file size
    size_kb = os.path.getsize(pdf_path) / 1024
    print(f"📄 File size: {size_kb:.1f} KB")

if __name__ == "__main__":
    docx_path = "outputs/reports/BP_Texas_City_Investigation_Report.docx"
    pdf_path = "outputs/reports/BP_Texas_City_Investigation_Report.pdf"
    
    print("🔄 Converting Word document to PDF...")
    print(f"📂 Source: {docx_path}")
    print(f"📂 Destination: {pdf_path}\n")
    
    try:
        convert_docx_to_pdf(docx_path, pdf_path)
        print("\n✅ Conversion completed successfully!")
    except Exception as e:
        print(f"\n❌ Error during conversion: {str(e)}")
        raise
